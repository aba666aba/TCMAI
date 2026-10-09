# Copyright (C) 2026 中医药智能问答系统 (TCMAI)
# SPDX-License-Identifier: AGPL-3.0-or-later
#
# 本文件是中医药智能问答系统 (TCMAI) 的一部分。
#
# 本程序是自由软件：你可以依据自由软件基金会发布的 GNU Affero 通用公共
# 许可证条款（第 3 版或你选择的任何更新版本）重新分发和/或修改本程序。
#
# 本程序按“希望有用”的原则分发，但不提供任何担保，甚至不提供适销性或
# 特定用途适用性的默示担保。详见 GNU Affero 通用公共许可证。
#
# 你应已随本程序收到 GNU Affero 通用公共许可证副本。如未收到，
# 请见 <https://www.gnu.org/licenses/>。
#
# 注意：以网络服务方式提供本程序的服务端修改版本，同样必须按 AGPL-3.0
# 第 13 条向所有用户提供对应源代码。
"""PPT小数位修复：修正第一轮替换中的误命中"""
from pptx import Presentation
from pathlib import Path

PPT = Path(r'E:\TOM\TCMAI\智慧中药方生成系统-结题答辩_modified.pptx')


def replace_in_para(para, old, new):
    full = para.text
    if old not in full:
        return False
    new_text = full.replace(old, new)
    ref_run = para.runs[0] if para.runs else None
    for run in para.runs:
        run.text = ''
    nr = para.add_run()
    nr.text = new_text
    if ref_run:
        nr.font.name = ref_run.font.name
        nr.font.size = ref_run.font.size
        nr.font.bold = ref_run.font.bold
        nr.font.italic = ref_run.font.italic
        if ref_run.font.color and ref_run.font.color.rgb:
            nr.font.color.rgb = ref_run.font.color.rgb
    return True


FIXES = [
    # S12: "提升至81.5%" was missed → 80.4%
    ('提升至81.5%，实现', '提升至80.4%，实现'),

    # S13 Shape6: 31.7% still there, 25.9% should be 26.7%
    ('31.7%降至25.9%', '31.1%降至26.7%'),
    ('仍有25.9%混淆率', '仍有26.7%混淆率'),

    # S14: 基线81.5% was wrongly changed to 80.4% → fix to 81.3%
    ('基线Recall@10达80.4%。KG', '基线Recall@10达81.3%。KG'),

    # S15: DL 74.8% → should be 76.3% (currently got 74.7% from S12 pattern)
    ('纯DL模式Recall@10达74.7%，Precision', '纯DL模式Recall@10达76.3%，Precision'),

    # S15: peak 81.5% → should be 82.5% (currently got 80.4% from S12 pattern)
    ('峰值达80.4%，Precision保持', '峰值达82.5%，Precision保持'),
]


def main():
    prs = Presentation(str(PPT))
    total = 0

    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for old, new in FIXES:
                    if replace_in_para(para, old, new):
                        print(f'  Slide {i}: "{old}" → "{new}"')
                        total += 1

    prs.save(str(PPT))
    print(f'\nDone: {total} fixes')


if __name__ == '__main__':
    main()
